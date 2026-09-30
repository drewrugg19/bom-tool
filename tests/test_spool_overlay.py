import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from openpyxl import load_workbook
from FabBOMTool.core import spool_overlay as overlay


def component(**values):
    row={'Fab Number':'SP-14','Mark':'01','Qty':'1','Size':'4"','Description':'PIPE',
         'Length':'10\'-0"','Way_Connector 1':'BW','Way_Connector 2':'BW','Heat ID':'H1'}
    row.update(values);return row


class OverlayTests(unittest.TestCase):
    def test_unchanged(self):
        row=overlay.compare_rows([component()],[component()])[0]
        self.assertEqual(row['Status'],'Unchanged');self.assertEqual(row['Changes'],'No Change')

    def test_every_compared_field(self):
        for field in overlay.FIELDS:
            with self.subTest(field=field):
                row=overlay.compare_rows([component()],[component(**{field:'Changed value'})])[0]
                self.assertEqual(row['Status'],'Modified')
                self.assertEqual(row['Old '+field],component()[field])
                self.assertEqual(row['New '+field],'Changed value')
                self.assertEqual(row['Changes'],f'{field}: {component()[field]} → Changed value')

    def test_multiple_changes_one_row(self):
        rows=overlay.compare_rows([component()],[component(Qty='2',Length='12\'-0"')])
        self.assertEqual(len(rows),1)
        self.assertEqual(rows[0]['Changes'],'Qty: 1 → 2; Length: 10\'-0" → 12\'-0"')

    def test_added_removed(self):
        added=overlay.compare_rows([],[component()])[0]
        removed=overlay.compare_rows([component()],[])[0]
        self.assertEqual(added['Status'],'Added');self.assertEqual(added['Old Fab Number'],'')
        self.assertEqual(removed['Status'],'Removed');self.assertEqual(removed['New Fab Number'],'')

    def test_renumbering_is_not_inferred(self):
        for field in ('Fab Number','Mark'):
            rows=overlay.compare_rows([component()],[component(**{field:'NEW'})])
            self.assertEqual({r['Status'] for r in rows},{'Added','Removed'})

    def test_exact_text_not_numeric_or_fuzzy_matching(self):
        row=overlay.compare_rows([component(Qty='1')],[component(Qty='1.0')])[0]
        self.assertEqual(row['Status'],'Modified');self.assertEqual(row['Old Qty'],'1')

    def test_sort_uses_new_or_old_fab_and_mark(self):
        old=[component(**{'Fab Number':'SP-2','Mark':'10'})]
        new=[component(**{'Fab Number':'SP-10'}),component(**{'Fab Number':'SP-2','Mark':'2'})]
        rows=overlay.compare_rows(old,new)
        self.assertEqual([(r['New Fab Number'] or r['Old Fab Number'],r['New Mark'] or r['Old Mark']) for r in rows], [('SP-2','2'),('SP-2','10'),('SP-10','01')])

    def test_missing_keys_not_matched(self):
        rows=overlay.compare_rows([component(Mark='')],[component(Mark='')])
        self.assertEqual({r['Status'] for r in rows},{'Added','Removed'})

    def test_conflicting_duplicates_not_guessed(self):
        rows=overlay.compare_rows([component(Qty='1'),component(Qty='2')],[component(Qty='3'),component(Qty='4')])
        self.assertEqual([r['Status'] for r in rows].count('Removed'),2)
        self.assertEqual([r['Status'] for r in rows].count('Added'),2)
        rows=overlay.compare_rows([component(Qty='1'),component(Qty='2')],[component(Qty='2'),component(Qty='1')])
        self.assertTrue(all(r['Status']=='Unchanged' for r in rows))

    def test_excel_columns_filters_literal_text_and_empty_export(self):
        expected=['Status','Old Fab Number','New Fab Number','Old Mark','New Mark','Old Qty','New Qty','Old Size','New Size','Old Description','New Description','Old Length','New Length','Old Way_Connector 1','New Way_Connector 1','Old Way_Connector 2','New Way_Connector 2','Old Heat ID','New Heat ID','Changes']
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/'delta.xlsx'
            overlay.export_overlay(overlay.compare_rows([],[component(Description='=1+1')]),path)
            wb=load_workbook(path);ws=wb.active
            self.assertEqual(wb.sheetnames,['Spool Overlay']);self.assertEqual([c.value for c in ws[1]],expected)
            self.assertEqual(ws.freeze_panes,'A2');self.assertEqual(ws.tables['SpoolOverlayResults'].autoFilter.ref,'A1:T2')
            self.assertEqual(ws.tables['SpoolOverlayResults'].tableStyleInfo.name,'TableStyleMedium9')
            self.assertEqual(ws['K2'].data_type,'s');wb.close()
            overlay.export_overlay([],path)
            wb=load_workbook(path);self.assertEqual(wb.active.max_row,2);wb.close()

    def test_multiple_real_pdfs_and_status_counts(self):
        with tempfile.TemporaryDirectory() as directory:
            paths=[Path(directory)/name for name in ('old1.pdf','old2.pdf','new1.pdf','new2.pdf')]
            write_pdf(paths[0]);write_pdf(paths[1],fab='SP-2')
            write_pdf(paths[2],pipe_length='12\'-0"');write_pdf(paths[3],fab='SP-2')
            result=overlay.run_overlay(paths[:2],paths[2:],'delta',directory)
            self.assertEqual(result['counts'],{'Added':0,'Removed':0,'Modified':1,'Unchanged':1})
            self.assertEqual(result['pdf_count'],4);self.assertEqual(result['err_rows'],0)
            changed=next(r for r in result['rows'] if r['Status']=='Modified')
            self.assertEqual(changed['Old Length'],'10\'-0"');self.assertEqual(changed['New Length'],'12\'-0"')
            second=overlay.run_overlay(paths[:2],paths[2:],'delta',directory)
            self.assertNotEqual(result['output_filename'],second['output_filename'])

    def test_parser_errors_are_diagnostics_not_material_rows(self):
        with tempfile.TemporaryDirectory() as directory, patch.object(overlay,'extract_spool_pdf',side_effect=ValueError('bad PDF')):
            result=overlay.run_overlay([Path('old.pdf')],[Path('new.pdf')],'delta',directory)
            self.assertEqual(result['rows'],[]);self.assertEqual(result['err_rows'],2)
            self.assertTrue(Path(result['output']).exists())

def write_pdf(path, fab='SP-14', pipe_length='10\'-0"'):
    """Tiny valid vector PDF fixture, exercised through the real pdfplumber parser."""
    def text(x,y,value):
        value=value.replace('\\','\\\\').replace('(','\\(').replace(')','\\)')
        return f'BT /F1 9 Tf {x} {y} Td ({value}) Tj ET'
    xs=[20,70,120,180,390,480,630,780,960]
    commands=['0.5 w']
    for x in xs: commands.append(f'{x} 670 m {x} 720 l S')
    for y in [670,695,720]: commands.append(f'20 {y} m 960 {y} l S')
    headers=['Mark','Qty','Size','Description','Length','Way_Connector 1','Way_Connector 2','Heat ID']
    values=['01','1','4"','PIPE stainless 316',pipe_length,'BW','BW','H101']
    commands += [text(x+3,705,v) for x,v in zip(xs,headers)]
    commands += [text(x+3,680,v) for x,v in zip(xs,values)]
    commands.append(text(550,100,'FAB NUMBER '+fab))
    stream='\n'.join(commands).encode('ascii')
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
             b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 1000 800] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
             b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>',
             b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream']
    data=b'%PDF-1.4\n';offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(data));data+=f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n'
    start=len(data);data+=b'xref\n0 6\n0000000000 65535 f \n'
    for offset in offsets[1:]:data+=f'{offset:010} 00000 n \n'.encode()
    data+=f'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode()
    Path(path).write_bytes(data)

