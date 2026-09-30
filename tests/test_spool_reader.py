import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import load_workbook

from FabBOMTool.core import spool_reader


class FakeCrop:
    def __init__(self, words=None, tables=None):
        self.words = words or []
        self.tables = tables or []

    def extract_words(self, **_kwargs):
        return self.words

    def extract_tables(self):
        return self.tables


class FakePage:
    width = 1000
    height = 800

    def __init__(self, words, tables):
        self.words = words
        self.tables = tables

    def crop(self, box):
        return FakeCrop(words=self.words) if box[1] >= self.height * 0.5 else FakeCrop(tables=self.tables)


class SpoolReaderTestCase(unittest.TestCase):
    def test_table_mapping_preserves_blank_heat_id(self):
        table = [
            ["Mark", "Qty", "Size", "Description", "Length", "Way_\nConnector 1", "Way Connector 2", "Heat ID"],
            ["10", "2", '4"', "Pipe", "12'-0\"", "BW", "FLG", None],
        ]
        self.assertEqual(
            spool_reader._table_rows(table),
            [{"Mark": "10", "Qty": "2", "Size": '4"', "Description": "Pipe", "Length": "12'-0\"", "Way_Connector 1": "BW", "Way_Connector 2": "FLG", "Heat ID": ""}],
        )

    def test_fab_number_is_read_from_bottom_right_title_block(self):
        words = [
            {"text": "FAB", "top": 610, "x0": 700},
            {"text": "NUMBER", "top": 610, "x0": 750},
            {"text": "B0018-S1-IP-A1-CHWR-SP-01", "top": 610, "x0": 830},
        ]
        self.assertEqual(spool_reader._fab_number(FakePage(words, [])), "B0018-S1-IP-A1-CHWR-SP-01")

    def test_multiple_files_are_combined_and_naturally_sorted(self):
        results = {
            "b.pdf": ([{"Fab Number": "B10", "Mark": "1", **{key: "" for key in spool_reader.SOURCE_COLUMNS if key != "Mark"}}], []),
            "a.pdf": ([{"Fab Number": "B2", "Mark": "10", **{key: "" for key in spool_reader.SOURCE_COLUMNS if key != "Mark"}}, {"Fab Number": "B2", "Mark": "2", **{key: "" for key in spool_reader.SOURCE_COLUMNS if key != "Mark"}}], []),
        }
        with tempfile.TemporaryDirectory() as temp_dir, mock.patch.object(
            spool_reader, "extract_spool_pdf", side_effect=lambda path: results[str(path)]
        ):
            result = spool_reader.run_spool_reader(["b.pdf", "a.pdf"], "spools", temp_dir)
            self.assertEqual([(row["Fab Number"], row["Mark"]) for row in result["rows"]], [("B2", "2"), ("B2", "10"), ("B10", "1")])

    def test_excel_has_exact_columns_and_filter_table_style(self):
        row = {column: "" for column in spool_reader.SPOOL_COLUMNS}
        row.update({"Fab Number": "B001", "Mark": "1", "Qty": "2"})
        with tempfile.TemporaryDirectory() as temp_dir:
            output = Path(temp_dir) / "spools.xlsx"
            spool_reader.export_spool_rows([row], output)
            workbook = load_workbook(output)
            self.assertEqual(workbook.sheetnames, ["Spool Drawings"])
            sheet = workbook.active
            self.assertEqual([cell.value for cell in sheet[1]], spool_reader.SPOOL_COLUMNS)
            self.assertEqual(list(sheet.tables), ["SpoolDrawingResults"])
            table = sheet.tables["SpoolDrawingResults"]
            self.assertEqual(table.ref, "A1:I2")
            self.assertEqual(table.tableStyleInfo.name, "TableStyleMedium9")
            self.assertNotIn("Inches", spool_reader.SPOOL_COLUMNS)
            self.assertNotIn("Multiplier", spool_reader.SPOOL_COLUMNS)


if __name__ == "__main__":
    unittest.main()
