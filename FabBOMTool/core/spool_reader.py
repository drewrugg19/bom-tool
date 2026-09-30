"""Extract material callouts and FAB numbers from spool drawing PDFs."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .logic import EXPORTS_DIR
from .export_files import reserve_export

SPOOL_COLUMNS = [
    "Fab Number", "Mark", "Qty", "Size", "Description", "Length",
    "Way_Connector 1", "Way_Connector 2", "Heat ID",
]
SOURCE_COLUMNS = SPOOL_COLUMNS[1:]


def _clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("\n", " ")).strip()


def _header_key(value: object) -> str:
    return re.sub(r"[^a-z0-9]", "", _clean(value).lower())


HEADER_ALIASES = {
    "mark": "Mark",
    "qty": "Qty",
    "quantity": "Qty",
    "size": "Size",
    "description": "Description",
    "length": "Length",
    "wayconnector1": "Way_Connector 1",
    "wayconnector2": "Way_Connector 2",
    "heatid": "Heat ID",
}


def _table_rows(table: list[list[object]]) -> list[dict[str, str]]:
    """Map a detected table only when it contains every requested header."""
    for header_index, raw_header in enumerate(table):
        mapped = [HEADER_ALIASES.get(_header_key(cell)) for cell in raw_header]
        if not all(column in mapped for column in SOURCE_COLUMNS):
            continue
        indexes = {column: mapped.index(column) for column in SOURCE_COLUMNS}
        rows = []
        for raw_row in table[header_index + 1:]:
            row = {
                column: _clean(raw_row[index]) if index < len(raw_row) else ""
                for column, index in indexes.items()
            }
            if any(row.values()):
                rows.append(row)
        return rows
    return []


def _fab_number(page) -> str:
    """Read FAB NUMBER from the bottom-right title block, not drawing geometry."""
    crop = page.crop((page.width * 0.5, page.height * 0.5, page.width, page.height))
    words = crop.extract_words(use_text_flow=True, keep_blank_chars=False) or []
    lines: dict[int, list[dict]] = {}
    for word in words:
        lines.setdefault(round(float(word["top"]) / 4), []).append(word)
    ordered = [
        " ".join(str(word["text"]) for word in sorted(line, key=lambda item: item["x0"]))
        for _, line in sorted(lines.items())
    ]
    text = "\n".join(ordered)
    match = re.search(r"FAB\s*NUMBER\s*[:#-]?\s*([A-Z0-9][A-Z0-9._/-]*)", text, re.I)
    if match:
        return match.group(1).strip()

    # Some title blocks put the value directly below the label.
    for index, line in enumerate(ordered):
        if re.search(r"\bFAB\s*NUMBER\b", line, re.I) and index + 1 < len(ordered):
            candidate = re.search(r"[A-Z0-9][A-Z0-9._/-]+", ordered[index + 1], re.I)
            if candidate:
                return candidate.group(0)
    return ""


def extract_spool_pdf(pdf_path: str | Path) -> tuple[list[dict[str, str]], list[str]]:
    rows: list[dict[str, str]] = []
    warnings: list[str] = []
    with pdfplumber.open(str(pdf_path)) as pdf:
        for page_number, page in enumerate(pdf.pages, start=1):
            fab_number = _fab_number(page)
            if not fab_number:
                warnings.append(f"Page {page_number}: FAB NUMBER not found")
            page_rows: list[dict[str, str]] = []
            top = page.crop((0, 0, page.width, page.height * 0.5))
            for table in top.extract_tables() or []:
                page_rows = _table_rows(table)
                if page_rows:
                    break
            if not page_rows:
                warnings.append(f"Page {page_number}: material/callout table not found")
            for row in page_rows:
                rows.append({"Fab Number": fab_number, **row})
    return rows, warnings


def _natural_key(value: str) -> list[object]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def export_spool_rows(rows: list[dict[str, str]], output_path: str | Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Spool Drawings"
    ws.append(SPOOL_COLUMNS)
    for row in rows:
        ws.append([row.get(column, "") for column in SPOOL_COLUMNS])
    ws.freeze_panes = "A2"
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for column_index, column in enumerate(SPOOL_COLUMNS, start=1):
        values = [column] + [str(row.get(column, "")) for row in rows]
        ws.column_dimensions[get_column_letter(column_index)].width = min(max(10, max(map(len, values)) + 2), 80)
    table = Table(displayName="SpoolDrawingResults", ref=f"A1:I{len(rows) + 1}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium9", showFirstColumn=False, showLastColumn=False,
        showRowStripes=True, showColumnStripes=False,
    )
    ws.add_table(table)
    wb.save(output_path)


def run_spool_reader(
    pdf_paths: Iterable[str | Path], export_filename: str, export_dir: str | Path | None = None,
) -> dict:
    combined: list[dict[str, str]] = []
    diagnostics = []
    for pdf_path in pdf_paths:
        try:
            rows, warnings = extract_spool_pdf(pdf_path)
            combined.extend(rows)
            diagnostics.append({"file": Path(pdf_path).name, "rows": len(rows), "warnings": warnings})
        except Exception as exc:
            diagnostics.append({"file": Path(pdf_path).name, "rows": 0, "error": f"{type(exc).__name__}: {exc}"})
    combined.sort(key=lambda row: (_natural_key(row["Fab Number"]), _natural_key(row["Mark"])))
    destination = Path(export_dir) if export_dir else EXPORTS_DIR
    destination.mkdir(parents=True, exist_ok=True)
    with reserve_export(destination, export_filename) as output:
        export_spool_rows(combined, output)
    return {
        "ok": True, "rows": combined, "row_count": len(combined),
        "diagnostics": diagnostics, "output": str(output), "output_filename": output.name,
    }
