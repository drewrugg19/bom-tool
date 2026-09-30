"""Exact Fab Number + Mark comparison of extracted material tables."""
from collections import Counter, defaultdict
from itertools import zip_longest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from .spool_reader import extract_spool_pdf, _natural_key
from .export_files import reserve_export

FIELDS = ('Qty', 'Size', 'Description', 'Length', 'Way_Connector 1', 'Way_Connector 2', 'Heat ID')
COLUMNS = ['Status', 'Old Fab Number', 'New Fab Number', 'Old Mark', 'New Mark'] + [
    f'{side} {field}' for field in FIELDS for side in ('Old', 'New')
] + ['Changes']


def value(row, field):
    item = row.get(field, '') if row is not None else ''
    return '' if item is None else str(item)


def key(row):
    return value(row, 'Fab Number'), value(row, 'Mark')


def compared(old, new):
    row = {f'{side} {field}': value(source, field)
           for field in ('Fab Number', 'Mark') + FIELDS for side, source in (('Old', old), ('New', new))}
    changes = [f'{field}: {value(old, field)} → {value(new, field)}'
               for field in FIELDS if value(old, field) != value(new, field)]
    status = 'Added' if old is None else 'Removed' if new is None else 'Modified' if changes else 'Unchanged'
    row.update(Status=status, Changes=status if status in ('Added', 'Removed') else '; '.join(changes) if changes else 'No Change')
    return row


def compare_rows(old, new):
    groups = [defaultdict(list), defaultdict(list)]
    output = []
    for rows, group, side in ((old, groups[0], 'Old'), (new, groups[1], 'New')):
        for row in rows:
            if all(key(row)):
                group[key(row)].append(row)
            else:
                # A missing key never identifies a row on the other side.
                output.append(compared(row if side == 'Old' else None, row if side == 'New' else None))
    for identity in groups[0].keys() | groups[1].keys():
        left, right = list(groups[0][identity]), list(groups[1][identity])
        # Cancel exact duplicates first, without relying on their PDF ordering.
        for a in left[:]:
            match = next((i for i, b in enumerate(right) if all(value(a, f) == value(b, f) for f in FIELDS)), None)
            if match is not None:
                output.append(compared(a, right.pop(match))); left.remove(a)
        if len(left) <= 1 and len(right) <= 1:
            output.extend(compared(a, b) for a, b in zip_longest(left, right))
        else:
            # Duplicate keys with conflicting values have no unique direct match.
            output.extend(compared(a, None) for a in left)
            output.extend(compared(None, b) for b in right)
    return sorted(output, key=lambda row: (
        _natural_key(row['New Fab Number'] or row['Old Fab Number']),
        _natural_key(row['New Mark'] or row['Old Mark']), row['Status']))


def export_overlay(rows, path):
    wb = Workbook(); ws = wb.active; ws.title = 'Spool Overlay'
    ws.append(COLUMNS)
    for row in rows:
        ws.append([row.get(column, '') for column in COLUMNS])
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str): cell.data_type = 's'
    if not rows: ws.append([''] * len(COLUMNS))
    ws.freeze_panes = 'A2'
    for cell in ws[1]:
        cell.font = Font(bold=True); cell.alignment = Alignment(horizontal='center')
    for index, column in enumerate(COLUMNS, 1):
        width = 65 if column == 'Changes' else 38 if 'Description' in column else 32 if 'Fab Number' in column else max(14, len(column) + 2)
        ws.column_dimensions[get_column_letter(index)].width = width
    table = Table(displayName='SpoolOverlayResults', ref=f'A1:T{ws.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium9', showRowStripes=True)
    ws.add_table(table); wb.save(path)


def run_overlay(old_paths, new_paths, export_filename, export_dir, source_names=None):
    diagnostics, sides = [], []
    for side, paths in [('OLD', old_paths), ('NEW', new_paths)]:
        rows = []
        for index, path in enumerate(paths):
            name = source_names[side.lower()][index] if source_names else Path(path).name
            try:
                extracted, warnings = extract_spool_pdf(path)
                warnings = list(warnings)
                if not extracted and not warnings: warnings.append('No material rows extracted.')
                if any(not all(key(row)) for row in extracted):
                    warnings.append('Missing Fab Number or Mark; these rows cannot be directly matched.')
                rows.extend(extracted)
                diagnostics.append({'side': side, 'file': name, 'warnings': warnings, 'rows': len(extracted)})
            except Exception as exc:
                diagnostics.append({'side': side, 'file': name, 'error': f'{type(exc).__name__}: {exc}'})
        duplicate_keys = [identity for identity, count in Counter(key(row) for row in rows).items() if count > 1]
        if duplicate_keys:
            diagnostics.append({'side': side, 'file': '', 'warnings': [
                'Duplicate Fab Number + Mark keys: ' + ', '.join(' / '.join(k) for k in duplicate_keys) +
                '. Identical rows are paired first; conflicting duplicates without a unique counterpart are reported as Added/Removed.']})
        sides.append(rows)
    rows = compare_rows(*sides)
    counts = {status: sum(row['Status'] == status for row in rows) for status in ('Added', 'Removed', 'Modified', 'Unchanged')}
    with reserve_export(export_dir, export_filename) as output:
        export_overlay(rows, output)
    return {'ok': True, 'rows': rows, 'row_count': len(rows), 'counts': counts,
            'output': str(output), 'output_filename': output.name, 'diagnostics': diagnostics,
            'warn_rows': sum(len(d.get('warnings', [])) for d in diagnostics),
            'err_rows': sum(bool(d.get('error')) for d in diagnostics), 'pdf_count': len(old_paths)+len(new_paths)}
